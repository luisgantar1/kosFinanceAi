"""
Modul pembuat file template Excel (.xlsx) interaktif untuk perencanaan
dan pencatatan keuangan anak kos menggunakan openpyxl.
"""

import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_budget_excel(
    total_allowance: float,
    kos_amount: float,
    kebutuhan_amount: float,
    makan_amount: float,
    hiburan_amount: float,
    investasi_amount: float,
    notes: str = "Disusun otomatis oleh KosFinance AI"
) -> io.BytesIO:
    """
    Menghasilkan file Excel (.xlsx) dengan 2 sheet:
    1. Ringkasan Anggaran & Evaluasi (dengan formula selisih)
    2. Pelacak Pengeluaran Harian (Daily Expense Tracker)
    """
    wb = Workbook()

    # Warna & Gaya Modern
    primary_color = "1E3A8A"   # Navy Blue
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=16, bold=True, color="1E3A8A")
    subtitle_font = Font(name="Calibri", size=10, italic=True, color="475569")
    bold_font = Font(name="Calibri", size=11, bold=True)
    normal_font = Font(name="Calibri", size=11)

    header_fill = PatternFill(start_color=primary_color, end_color=primary_color, fill_type="solid")
    accent_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    total_fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")

    thin_border_side = Side(style='thin', color='CBD5E1')
    table_border = Border(
        left=thin_border_side, right=thin_border_side,
        top=thin_border_side, bottom=thin_border_side
    )
    double_bottom_border = Border(
        left=thin_border_side, right=thin_border_side,
        top=thin_border_side, bottom=Side(style='double', color='0F172A')
    )

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    # ==========================================
    # SHEET 1: Rencana Anggaran Bulanan
    # ==========================================
    ws1 = wb.active
    ws1.title = "Rencana Anggaran Bulanan"
    ws1.views.sheetView[0].showGridLines = True

    # Judul
    ws1.merge_cells("A1:G1")
    ws1["A1"] = "RENCANA ANGGARAN & ARUS KAS BULANAN ANAK KOS"
    ws1["A1"].font = title_font
    ws1["A1"].alignment = align_left

    ws1.merge_cells("A2:G2")
    ws1["A2"] = f"Panduan Resmi KosFinance AI | Catatan: {notes}"
    ws1["A2"].font = subtitle_font
    ws1["A2"].alignment = align_left

    # Info Ringkasan Uang Saku
    ws1["A4"] = "Total Uang Saku Bulanan:"
    ws1["A4"].font = bold_font
    ws1["B4"] = total_allowance
    ws1["B4"].font = bold_font
    ws1["B4"].number_format = '"Rp"#,##0'

    # Header Tabel
    headers = [
        "No", "Kategori Pos Pengeluaran", "Porsi (%)", 
        "Anggaran Alokasi (Rp)", "Pagu / Keterangan", 
        "Realisasi Aktual (Rp)", "Selisih (+/-)"
    ]

    header_row = 6
    for col_idx, header in enumerate(headers, 1):
        cell = ws1.cell(row=header_row, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_center
        cell.border = table_border

    # Data Baris
    items = [
        (1, "Sewa Kos & Utilitas (Listrik/Air)", kos_amount, "Biaya sewa kamar & utilitas tetap"),
        (2, "Kebutuhan Pokok & Operasional Kos", kebutuhan_amount, "Sanitasi, galon, laundry, bensin"),
        (3, "Makan & Minum Harian", makan_amount, f"Pagu harian: Rp {int(makan_amount/30):,}/hari"),
        (4, "Senang-senang / Gaya Hidup", hiburan_amount, f"Pagu mingguan: Rp {int(hiburan_amount/4):,}/minggu"),
        (5, "Investasi, Tabungan & Dana Darurat", investasi_amount, "RDPU, emas mini, dana darurat likuid"),
    ]

    for idx, (no, nama_pos, nominal, pagu) in enumerate(items, start=7):
        ws1.cell(row=idx, column=1, value=no).alignment = align_center
        ws1.cell(row=idx, column=2, value=nama_pos).alignment = align_left
        
        # Porsi (%) via formula
        persen_cell = ws1.cell(row=idx, column=3, value=f"=D{idx}/B4")
        persen_cell.number_format = '0.0%'
        persen_cell.alignment = align_center

        # Nominal Alokasi
        nom_cell = ws1.cell(row=idx, column=4, value=nominal)
        nom_cell.number_format = '"Rp"#,##0'
        nom_cell.alignment = align_right

        # Pagu / Keterangan
        ws1.cell(row=idx, column=5, value=pagu).alignment = align_left

        # Realisasi (Diisi User)
        realisasi_cell = ws1.cell(row=idx, column=6)
        realisasi_cell.number_format = '"Rp"#,##0'
        realisasi_cell.alignment = align_right

        # Selisih (+/-) via formula: Anggaran - Realisasi
        selisih_cell = ws1.cell(row=idx, column=7, value=f"=D{idx}-F{idx}")
        selisih_cell.number_format = '"Rp"#,##0'
        selisih_cell.alignment = align_right

        # Styling borders & fonts
        for c in range(1, 8):
            cell = ws1.cell(row=idx, column=c)
            cell.font = normal_font
            cell.border = table_border
            if idx % 2 == 0:
                cell.fill = accent_fill

    # Total Row
    tot_row = 12
    ws1.cell(row=tot_row, column=1, value="").border = double_bottom_border
    ws1.cell(row=tot_row, column=2, value="TOTAL ANGGARAN").alignment = align_left
    ws1.cell(row=tot_row, column=2).font = bold_font
    
    # Formula total persen
    tot_pct = ws1.cell(row=tot_row, column=3, value=f"=SUM(C7:C11)")
    tot_pct.number_format = '0.0%'
    tot_pct.alignment = align_center
    tot_pct.font = bold_font

    # Formula total nominal
    tot_nom = ws1.cell(row=tot_row, column=4, value=f"=SUM(D7:D11)")
    tot_nom.number_format = '"Rp"#,##0'
    tot_nom.alignment = align_right
    tot_nom.font = bold_font

    ws1.cell(row=tot_row, column=5, value="100% Terencana Bebas Defisit").alignment = align_center
    ws1.cell(row=tot_row, column=5).font = bold_font

    # Total realisasi formula
    tot_real = ws1.cell(row=tot_row, column=6, value=f"=SUM(F7:F11)")
    tot_real.number_format = '"Rp"#,##0'
    tot_real.alignment = align_right
    tot_real.font = bold_font

    # Total selisih formula
    tot_sel = ws1.cell(row=tot_row, column=7, value=f"=SUM(G7:G11)")
    tot_sel.number_format = '"Rp"#,##0'
    tot_sel.alignment = align_right
    tot_sel.font = bold_font

    for c in range(1, 8):
        cell = ws1.cell(row=tot_row, column=c)
        cell.border = double_bottom_border
        cell.fill = total_fill

    # ==========================================
    # SHEET 2: Pelacak Pengeluaran Harian
    # ==========================================
    ws2 = wb.create_sheet(title="Log Pengeluaran Harian")
    ws2.views.sheetView[0].showGridLines = True

    ws2.merge_cells("A1:E1")
    ws2["A1"] = "CATATAN TRANSAKSI PENGELUARAN HARIAN MAHASISWA"
    ws2["A1"].font = title_font

    ws2.merge_cells("A2:E2")
    ws2["A2"] = "Isi setiap pengeluaran harian Anda di sini untuk memantau pagu pengeluaran riil."
    ws2["A2"].font = subtitle_font

    tracker_headers = ["Tanggal", "Kategori Pos", "Rincian Transaksi / Barang", "Nominal (Rp)", "Metode Pembayaran"]
    for col_idx, h in enumerate(tracker_headers, 1):
        cell = ws2.cell(row=4, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_center
        cell.border = table_border

    # Baris contoh
    sample_rows = [
        ("01/10/2026", "Sewa Kos & Utilitas", "Bayar sewa kamar kos bulan ini", kos_amount, "Transfer Bank"),
        ("02/10/2026", "Makan & Minum Harian", "Makan siang warteg + es teh", 18000, "Tunai / QRIS"),
        ("02/10/2026", "Kebutuhan Pokok", "Isi ulang galon air minum", 7000, "Tunai"),
        ("03/10/2026", "Senang-senang / Gaya Hidup", "Nongkrong warkop nugas tugas", 25000, "QRIS"),
    ]

    for r_idx, row_data in enumerate(sample_rows, start=5):
        for c_idx, val in enumerate(row_data, start=1):
            cell = ws2.cell(row=r_idx, column=c_idx, value=val)
            cell.font = normal_font
            cell.border = table_border
            if c_idx == 4:
                cell.number_format = '"Rp"#,##0'
                cell.alignment = align_right
            elif c_idx == 1:
                cell.alignment = align_center
            else:
                cell.alignment = align_left

    # Auto-adjust column width for both sheets
    for ws in [ws1, ws2]:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or '')
                if not cell.coordinate in ws.merged_cells:
                    max_len = max(max_len, len(val_str))
            ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    # Simpan ke dalam buffer memori BytesIO
    output_stream = io.BytesIO()
    wb.save(output_stream)
    output_stream.seek(0)
    return output_stream
